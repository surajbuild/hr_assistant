"""
app/rag/document_loader.py
--------------------------
Text extraction for uploaded HR documents (PRD Sections 12 & 23).

Supported formats:
- PDF  : pypdf, one entry per page (page numbers are 1-based) so answers can cite pages.
- DOCX : python-docx paragraphs + table cells (no page concept → page=None).
- TXT  : UTF-8 (falls back to latin-1) plain text (page=None).

Returns a list of (page_number | None, text) tuples.
"""

from pathlib import Path
from typing import List, Optional, Tuple

SUPPORTED_EXTENSIONS = {"pdf", "docx", "txt"}

PageText = Tuple[Optional[int], str]


class DocumentLoadError(Exception):
    """Raised when a document cannot be parsed."""
    pass


def _load_pdf(path: Path) -> List[PageText]:
    from pypdf import PdfReader

    try:
        reader = PdfReader(str(path))
    except Exception as exc:  # corrupted / encrypted files
        raise DocumentLoadError(f"Could not read PDF: {exc}") from exc
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:
            raise DocumentLoadError("PDF is password protected.") from exc
    pages: List[PageText] = []
    for index, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        pages.append((index, text))
    return pages


def _load_docx(path: Path) -> List[PageText]:
    import docx

    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise DocumentLoadError(f"Could not read DOCX: {exc}") from exc
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return [(None, "\n".join(parts))]


def _load_txt(path: Path) -> List[PageText]:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return [(None, raw.decode(encoding))]
        except UnicodeDecodeError:
            continue
    raise DocumentLoadError("Could not decode text file.")


def load_document(path: str, file_type: str) -> List[PageText]:
    """
    Extract text from a stored document.

    Raises:
        DocumentLoadError: unsupported type, unreadable file, or no extractable text.
    """
    file_type = file_type.lower().lstrip(".")
    file_path = Path(path)
    if file_type not in SUPPORTED_EXTENSIONS:
        raise DocumentLoadError(f"Unsupported file type '.{file_type}'. Allowed: PDF, DOCX, TXT.")
    if not file_path.exists():
        raise DocumentLoadError("Stored file not found.")

    loader = {"pdf": _load_pdf, "docx": _load_docx, "txt": _load_txt}[file_type]
    pages = loader(file_path)
    if not any(text.strip() for _, text in pages):
        raise DocumentLoadError(
            "No extractable text found (scanned/image-only documents are not supported)."
        )
    return pages
