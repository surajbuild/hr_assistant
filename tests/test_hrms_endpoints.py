"""
tests/test_hrms_endpoints.py
----------------------------
Integration tests for the HRMS endpoints added on 2026-10-05:

  /auth/me, employees CRUD + manager scope, /departments,
  attendance today/check-in/check-out/daily/records + calculation engine,
  leaves list/balance/cancel + manager team-only approvals,
  /salary payroll sheet, /users (admin), /chat/history, /chat/logs,
  /dashboard/me, /dashboard/summary monthly trend, /reports/* period filters.

Runs against the real MySQL database from .env and assumes `python scripts/seed_db.py`
has been run (6 seeded employees: Vikram=admin, Neha=hr, Priya=manager of Aman & Rahul,
Sneha reports to Vikram). Test rows it creates are removed at the end.
"""

import io
import os
import sys
from datetime import date, datetime, time

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
from fastapi.testclient import TestClient

from app.database.connection import SessionLocal
from app.database.models import Attendance, Employee, Leave, User
from app.main import app
from app.services.attendance_service import calculate_day_metrics, check_in, check_out
from app.services.leave_service import count_leave_days
from app.utils.security import create_access_token

client = TestClient(app)
db = SessionLocal()

passed_count = 0
failed_count = 0
TEST_CODE = "T-HRMS-01"
TEST_EMAIL = "t.hrms01@hrtest.dev"


def chk(condition: bool, msg: str, fail_detail: str = ""):
    global passed_count, failed_count
    if condition:
        print(f"  [PASS] {msg}")
        passed_count += 1
    else:
        print(f"  [FAIL] {msg} -> {fail_detail}")
        failed_count += 1


def headers_for(email: str):
    user = db.query(User).filter(User.email == email).first()
    if not user:
        raise RuntimeError(f"Seed user {email} missing — run scripts/seed_db.py")
    return user, {"Authorization": f"Bearer {create_access_token(user_id=user.id, role=user.role)}"}


def emp_id(code: str) -> int:
    return db.query(Employee).filter(Employee.employee_code == code).first().id


def cleanup():
    db.rollback()
    emp = db.query(Employee).filter(Employee.employee_code == TEST_CODE).first()
    if emp:
        db.query(Attendance).filter(Attendance.employee_id == emp.id).delete()
        db.query(Leave).filter(Leave.employee_id == emp.id).delete()
        db.query(User).filter(User.employee_id == emp.id).delete()
        db.delete(emp)
    db.query(Leave).filter(Leave.reason == "T-HRMS test leave").delete()
    db.commit()


try:
    cleanup()
    admin_u, admin_h = headers_for("admin@company.com")
    hr_u, hr_h = headers_for("neha.hr@company.com")
    mgr_u, mgr_h = headers_for("priya.mgr@company.com")
    aman_u, aman_h = headers_for("aman@company.com")
    rahul_u, rahul_h = headers_for("rahul@company.com")
    AMAN, RAHUL, SNEHA, PRIYA = emp_id("EMP004"), emp_id("EMP005"), emp_id("EMP006"), emp_id("EMP003")

    # ------------------------------------------------------------------
    print("\n[1] GET /auth/me")
    r = client.get("/auth/me", headers=aman_h)
    chk(r.status_code == 200 and r.json()["role"] == "employee" and r.json()["employee"]["employee_code"] == "EMP004",
        "Employee identity + profile returned", r.text[:200])
    chk(r.json()["employee"]["manager_name"] == "Priya Nair", "Manager name included", r.text[:200])
    chk(client.get("/auth/me").status_code == 401, "Unauthenticated -> 401")

    # ------------------------------------------------------------------
    print("\n[2] Employee directory scope")
    r = client.get("/employees", headers=hr_h)
    chk(r.status_code == 200 and len(r.json()) >= 6, "HR sees all employees", str(r.status_code))
    chk(all("manager_name" in e and "email" in e for e in r.json()), "Directory rows include manager_name + email")
    r = client.get("/employees", headers=mgr_h)
    ids = {e["id"] for e in r.json()} if r.status_code == 200 else set()
    chk(r.status_code == 200 and ids == {PRIYA, AMAN, RAHUL}, "Manager sees only self + direct reports", str(ids))
    chk(client.get("/employees", headers=aman_h).status_code == 403, "Employee blocked from directory -> 403")
    r = client.get("/employees?search=aman", headers=hr_h)
    chk(r.status_code == 200 and [e["employee_code"] for e in r.json()] == ["EMP004"], "Search filter works", r.text[:200])
    r = client.get("/employees?department=Engineering", headers=hr_h)
    chk(r.status_code == 200 and all(e["department"] == "Engineering" for e in r.json()), "Department filter works")

    print("\n[3] GET /employees/{id} ownership")
    chk(client.get(f"/employees/{AMAN}", headers=aman_h).status_code == 200, "Employee can view own profile by id")
    chk(client.get(f"/employees/{RAHUL}", headers=aman_h).status_code == 403, "Employee cannot view a peer -> 403")
    chk(client.get(f"/employees/{AMAN}", headers=mgr_h).status_code == 200, "Manager can view a direct report")
    chk(client.get(f"/employees/{SNEHA}", headers=mgr_h).status_code == 403, "Manager cannot view non-team employee -> 403")
    chk(client.get("/employees/999999", headers=hr_h).status_code == 404, "HR missing employee -> 404")

    # ------------------------------------------------------------------
    print("\n[4] Employee create / update / soft delete")
    payload = {
        "employee_code": TEST_CODE, "name": "Test Hrms", "department": "QA", "designation": "QA Engineer",
        "joining_date": "2026-01-05", "manager_id": PRIYA,
        "email": TEST_EMAIL, "password": "TestPass#2026", "role": "employee",
    }
    chk(client.post("/employees", json=payload, headers=aman_h).status_code == 403, "Employee cannot create -> 403")
    chk(client.post("/employees", json=payload, headers=mgr_h).status_code == 403, "Manager cannot create -> 403")
    chk(client.post("/employees", json={**payload, "role": "admin", "employee_code": "T-HRMS-X", "email": "t.hrms.x@hrtest.dev"},
                    headers=hr_h).status_code == 403, "HR cannot mint admin accounts -> 403")
    r = client.post("/employees", json=payload, headers=hr_h)
    chk(r.status_code == 201 and r.json()["email"] == TEST_EMAIL and r.json()["manager_name"] == "Priya Nair",
        "HR creates employee with login", r.text[:200])
    new_id = r.json().get("id") if r.status_code == 201 else None
    chk(client.post("/employees", json=payload, headers=hr_h).status_code == 409, "Duplicate employee code -> 409")
    chk(client.post("/employees", json={**payload, "employee_code": "T-HRMS-02", "password": None}, headers=hr_h).status_code in (400, 409),
        "Login without password rejected")
    login = client.post("/auth/login", json={"email": TEST_EMAIL, "password": "TestPass#2026"})
    chk(login.status_code == 200, "New login account can sign in", login.text[:120])
    new_h = {"Authorization": f"Bearer {login.json().get('access_token', '')}"}

    r = client.put(f"/employees/{new_id}", json={"designation": "Senior QA Engineer"}, headers=hr_h)
    chk(r.status_code == 200 and r.json()["designation"] == "Senior QA Engineer", "HR partial update", r.text[:200])
    chk(client.put(f"/employees/{new_id}", json={"manager_id": new_id}, headers=hr_h).status_code == 400,
        "Self-manager rejected -> 400")
    chk(client.put(f"/employees/{new_id}", json={"designation": "X"}, headers=mgr_h).status_code == 403,
        "Manager cannot update -> 403")

    # ------------------------------------------------------------------
    print("\n[5] Attendance calculation engine")
    m = calculate_day_metrics(time(9, 0), time(18, 0))
    chk(m == {"working_minutes": 480, "overtime_minutes": 0, "late_minutes": 0}, "09:00-18:00 = 480 worked, 0 OT, 0 late", str(m))
    m = calculate_day_metrics(time(9, 25), time(20, 0))
    chk(m == {"working_minutes": 575, "overtime_minutes": 95, "late_minutes": 25}, "09:25-20:00 = 575 worked, 95 OT, 25 late", str(m))
    m = calculate_day_metrics(time(9, 10), time(13, 0))
    chk(m["late_minutes"] == 0 and m["working_minutes"] == 230, "Within grace = not late; short span keeps lunch", str(m))

    print("\n[6] Check-in / check-out")
    r = client.post("/attendance/check-in", headers=new_h)
    chk(r.status_code == 201 and r.json()["in_time"] is not None, "Check-in creates today's record", r.text[:200])
    chk(client.post("/attendance/check-in", headers=new_h).status_code == 409, "Second check-in -> 409")
    r = client.get("/attendance/today", headers=new_h)
    chk(r.status_code == 200 and r.json() is not None and r.json()["out_time"] is None, "GET /attendance/today returns open record")
    r = client.post("/attendance/check-out", headers=new_h)
    chk(r.status_code == 200 and r.json()["out_time"] is not None, "Check-out closes the record", r.text[:200])
    chk(client.post("/attendance/check-out", headers=new_h).status_code == 409, "Second check-out -> 409")
    # Deterministic service-level check of a full day
    db.query(Attendance).filter(Attendance.employee_id == new_id).delete()
    db.commit()
    rec = check_in(db, new_id, now=datetime(2026, 1, 6, 9, 30))
    rec = check_out(db, new_id, now=datetime(2026, 1, 6, 19, 30))
    chk(rec.late_minutes == 30 and rec.working_minutes == 540 and rec.overtime_minutes == 60 and rec.status == "present",
        "Service: 09:30-19:30 -> late 30, worked 540, OT 60", f"{rec.late_minutes}/{rec.working_minutes}/{rec.overtime_minutes}")

    print("\n[7] Daily sheet & record search")
    r = client.get("/attendance/daily?date=2024-08-02", headers=hr_h)
    chk(r.status_code == 200 and r.json()["counts"]["total"] >= 6, "HR daily sheet covers all active employees", r.text[:200])
    aman_row = next((x for x in r.json()["rows"] if x["employee_id"] == AMAN), None)
    chk(aman_row is not None and aman_row["late_minutes"] == 25, "Aman late 25 min on 2024-08-02", str(aman_row))
    r = client.get("/attendance/daily", headers=mgr_h)
    chk(r.status_code == 200 and {x["employee_id"] for x in r.json()["rows"]} <= {PRIYA, AMAN, RAHUL, new_id},
        "Manager daily sheet limited to team")
    chk(client.get("/attendance/daily", headers=aman_h).status_code == 403, "Employee blocked from daily sheet -> 403")
    r = client.get(f"/attendance/records?employee_id={AMAN}&from_date=2024-08-01&to_date=2024-08-31&status=present", headers=hr_h)
    chk(r.status_code == 200 and len(r.json()) == 22, "Aman: 22 present records in Aug 2024", str(len(r.json())))
    chk(client.get(f"/attendance/records?employee_id={SNEHA}", headers=mgr_h).status_code == 403,
        "Manager cannot search non-team records -> 403")
    chk(client.get("/attendance/records?from_date=2024-09-30&to_date=2024-09-01", headers=hr_h).status_code == 422,
        "Inverted date range -> 422")

    # ------------------------------------------------------------------
    print("\n[8] Leaves: day counting, balance, list scope")
    chk(count_leave_days(date(2024, 8, 12), date(2024, 8, 13)) == 2, "Mon-Tue = 2 leave days")
    chk(count_leave_days(date(2024, 8, 16), date(2024, 8, 19)) == 2, "Fri-Mon skips weekend = 2 days")
    r = client.get("/leaves/balance/me?year=2024", headers=aman_h)
    casual = next((b for b in r.json() if b["leave_type"] == "casual"), {}) if r.status_code == 200 else {}
    chk(casual.get("entitled") == 12 and casual.get("used") == 2 and casual.get("remaining") == 10,
        "Aman 2024 casual: 12 entitled, 2 used, 10 remaining", str(casual))
    # AI chat LEAVE tool includes the Python-calculated balance (was missing: "not available in context")
    from app.ai.router import Intent, retrieve_hr_context
    ctx, _src, _deny = retrieve_hr_context(db, aman_u, Intent.LEAVE, "How many casual leaves are remaining in 2024?")
    chk("Casual leave: entitled 12, used 2, pending approval 0, remaining 10" in ctx,
        "Chat leave context carries the 2024 casual balance", ctx[:300])
    chk("Working days: 2" in ctx, "Chat leave history uses working days (KI-020)", ctx[:600])
    r = client.get("/leaves", headers=mgr_h)
    chk(r.status_code == 200 and all(x["employee_id"] in {PRIYA, AMAN, RAHUL, new_id} for x in r.json()),
        "Manager leave list limited to team")
    r = client.get("/leaves?status=approved", headers=hr_h)
    chk(r.status_code == 200 and len(r.json()) > 0 and all(x["status"] == "approved" for x in r.json()), "HR status filter")
    chk(client.get("/leaves", headers=aman_h).status_code == 403, "Employee blocked from leave list -> 403")

    print("\n[9] Leave approvals: manager team-only, no self-approval")
    team_leave = Leave(employee_id=AMAN, leave_type="casual", from_date=date(2026, 11, 2), to_date=date(2026, 11, 3),
                       status="pending", reason="T-HRMS test leave")
    other_leave = Leave(employee_id=SNEHA, leave_type="casual", from_date=date(2026, 11, 2), to_date=date(2026, 11, 2),
                        status="pending", reason="T-HRMS test leave")
    own_leave = Leave(employee_id=PRIYA, leave_type="casual", from_date=date(2026, 11, 9), to_date=date(2026, 11, 9),
                      status="pending", reason="T-HRMS test leave")
    db.add_all([team_leave, other_leave, own_leave])
    db.commit()
    chk(client.patch(f"/leaves/{other_leave.id}/status", json={"status": "approved"}, headers=mgr_h).status_code == 403,
        "Manager cannot approve non-team leave -> 403")
    chk(client.patch(f"/leaves/{own_leave.id}/status", json={"status": "approved"}, headers=mgr_h).status_code == 403,
        "Manager cannot approve own leave -> 403")
    chk(client.patch(f"/leaves/{team_leave.id}/status", json={"status": "approved"}, headers=mgr_h).status_code == 200,
        "Manager approves direct report's leave")
    chk(client.patch(f"/leaves/{other_leave.id}/status", json={"status": "rejected"}, headers=hr_h).status_code == 200,
        "HR can act on any leave")

    print("\n[10] Cancel own pending leave")
    r = client.post("/leaves", json={"leave_type": "sick", "start_date": "2026-12-01", "end_date": "2026-12-01",
                                      "reason": "T-HRMS test leave"}, headers=aman_h)
    lid = r.json().get("id")
    chk(client.post(f"/leaves/{lid}/cancel", headers=rahul_h).status_code == 404, "Peer cannot cancel -> 404")
    r = client.post(f"/leaves/{lid}/cancel", headers=aman_h)
    chk(r.status_code == 200 and r.json()["status"] == "cancelled", "Owner cancels pending leave", r.text[:200])
    chk(client.post(f"/leaves/{lid}/cancel", headers=aman_h).status_code == 400, "Cancel twice -> 400")
    chk(client.post(f"/leaves/{team_leave.id}/cancel", headers=aman_h).status_code == 400, "Cannot cancel approved leave -> 400")

    # ------------------------------------------------------------------
    print("\n[11] Payroll sheet")
    # Expected default = latest (year, month) present in the salary table — data-independent
    from app.database.models import Salary
    db.commit()
    latest = db.query(Salary.year, Salary.month).order_by(Salary.year.desc(), Salary.month.desc()).first()
    r = client.get("/salary", headers=hr_h)
    chk(r.status_code == 200 and (r.json()["year"], r.json()["month"]) == (latest[0], latest[1]) and len(r.json()["items"]) > 0,
        f"Defaults to latest payroll month ({latest[1]}/{latest[0]})", r.text[:200])
    chk(all("employee_name" in i for i in r.json()["items"]), "Rows include employee names")
    r = client.get("/salary?month=8&year=2024", headers=admin_h)
    chk(r.status_code == 200 and all(i["month"] == 8 for i in r.json()["items"]), "Explicit month filter")
    chk(client.get("/salary?month=13&year=2024", headers=hr_h).status_code in (400, 422), "Invalid month rejected")
    chk(client.get("/salary", headers=mgr_h).status_code == 403, "Manager blocked from payroll -> 403")
    chk(client.get("/salary", headers=aman_h).status_code == 403, "Employee blocked from payroll -> 403")

    # ------------------------------------------------------------------
    print("\n[12] Departments")
    r = client.get("/departments", headers=hr_h)
    eng = next((d for d in r.json() if d["name"] == "Engineering"), None) if r.status_code == 200 else None
    chk(eng is not None and eng["employee_count"] >= 3 and "Priya Nair" in eng["managers"], "Engineering stats", str(eng))
    chk(client.get("/departments", headers=mgr_h).status_code == 200, "Manager can view departments")
    chk(client.get("/departments", headers=aman_h).status_code == 403, "Employee blocked -> 403")

    # ------------------------------------------------------------------
    print("\n[13] Users admin")
    r = client.get("/users", headers=admin_h)
    chk(r.status_code == 200 and any(u["email"] == "aman@company.com" for u in r.json()), "Admin lists users")
    chk(client.get("/users", headers=hr_h).status_code == 403, "HR blocked from user admin -> 403")
    chk(client.patch(f"/users/{admin_u.id}", json={"role": "employee"}, headers=admin_h).status_code == 400,
        "Admin cannot demote self -> 400")
    new_user = db.query(User).filter(User.email == TEST_EMAIL).first()
    r = client.patch(f"/users/{new_user.id}", json={"role": "manager"}, headers=admin_h)
    chk(r.status_code == 200 and r.json()["role"] == "manager", "Admin changes another user's role")

    # ------------------------------------------------------------------
    print("\n[14] Chat history & audit log")
    chk(client.get("/chat/history", headers=aman_h).status_code == 200, "Own chat history")
    r = client.get("/chat/logs?limit=5", headers=admin_h)
    chk(r.status_code == 200 and len(r.json()) <= 5, "Admin audit log")
    chk(client.get("/chat/logs", headers=hr_h).status_code == 403, "HR blocked from audit log -> 403")

    # ------------------------------------------------------------------
    print("\n[15] Dashboards")
    r = client.get("/dashboard/me", headers=aman_h)
    d = r.json() if r.status_code == 200 else {}
    chk(r.status_code == 200 and d["attendance"]["present_days"] > 0 and len(d["leave_balance"]) == 3 and d["team"] is None,
        "Employee personal dashboard", r.text[:200])
    r = client.get("/dashboard/me", headers=mgr_h)
    chk(r.status_code == 200 and r.json()["team"] is not None and r.json()["team"]["size"] >= 2, "Manager gets team snapshot")
    r = client.get("/dashboard/summary?date=2024-09-30", headers=hr_h)
    chk(r.status_code == 200 and len(r.json().get("monthly_attendance", [])) >= 2, "Summary includes monthly trend")
    aug = next((m for m in r.json()["monthly_attendance"] if m["month"] == "2024-08"), {})
    chk(aug.get("present") == 22 and aug.get("absent") == 2, "Aug 2024 trend = 22 present / 2 absent", str(aug))

    # ------------------------------------------------------------------
    print("\n[16] Reports with period filters")
    r = client.get("/reports/attendance?month=8&year=2024", headers=hr_h)
    wb = openpyxl.load_workbook(io.BytesIO(r.content)) if r.status_code == 200 else None
    chk(wb is not None and wb.sheetnames[0] == "Summary", "Attendance report opens on Summary sheet")
    if wb:
        header = [c.value for c in wb["Summary"][4]]
        chk(header[3:] == ["Present", "Absent", "Half Day", "Late Count", "Working Minutes", "OT Minutes"],
            "Attendance summary has PRD columns", str(header))
        aman_row = next((row for row in wb["Summary"].iter_rows(min_row=5, values_only=True) if row[0] == "EMP004"), None)
        chk(aman_row is not None and aman_row[3] == 22 and aman_row[4] == 2 and aman_row[6] == 4,
            "Aman Aug: 22 present, 2 absent, 4 late", str(aman_row))
    r = client.get("/reports/overtime?month=9&year=2024", headers=hr_h)
    wb = openpyxl.load_workbook(io.BytesIO(r.content)) if r.status_code == 200 else None
    rahul_row = next((row for row in wb["Summary"].iter_rows(min_row=5, values_only=True) if row and row[0] == "EMP005"), None) if wb else None
    chk(rahul_row is not None and rahul_row[3] == round(1115 / 60, 2), "Rahul Sep OT = 18.58 h", str(rahul_row))
    r = client.get("/reports/leave?year=2024", headers=hr_h)
    chk(r.status_code == 200 and "leave_report_2024" in r.headers.get("content-disposition", ""), "Leave report download")
    chk(client.get("/reports/leave", headers=aman_h).status_code == 403, "Employee blocked from reports -> 403")
    chk(client.get("/reports/leave?month=8", headers=hr_h).status_code == 422, "Month without year -> 422")

    # ------------------------------------------------------------------
    print("\n[17] Soft delete")
    chk(client.delete(f"/employees/{hr_u.employee_id}", headers=hr_h).status_code == 400, "Cannot deactivate yourself -> 400")
    r = client.delete(f"/employees/{new_id}", headers=hr_h)
    chk(r.status_code == 200 and r.json()["status"] == "inactive", "Soft delete marks employee inactive")
    db.commit()
    chk(db.query(User).filter(User.email == TEST_EMAIL).first().status == "inactive", "Linked login disabled")
    chk(client.post("/auth/login", json={"email": TEST_EMAIL, "password": "TestPass#2026"}).status_code == 403,
        "Deactivated user cannot log in")
    chk(db.query(Attendance).filter(Attendance.employee_id == new_id).count() > 0, "History preserved after soft delete")

except Exception as exc:  # report unexpected crashes as failures
    import traceback
    traceback.print_exc()
    chk(False, "Unexpected exception", str(exc))
finally:
    cleanup()
    db.close()

print("\n" + "=" * 65)
print(f"  HRMS ENDPOINT TEST RESULTS: {passed_count} PASSED, {failed_count} FAILED")
print("=" * 65)
sys.exit(1 if failed_count else 0)
