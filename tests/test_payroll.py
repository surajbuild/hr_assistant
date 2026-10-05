"""
tests/test_payroll.py
---------------------
Payroll engine (salary_service.generate_payroll, POST /salary/generate) and the
pro-rated overtime amount in the overtime report (KI-015). Rules: D-021.

Uses only test-owned employees (codes T-PAY-*, department "QA Payroll") in
June 2025 — 21 working days, no company holidays — and removes everything it
creates. Seeded rows are checked to be untouched.

Hand-checked example (employee T-PAY-EMP, gross 84,000):
    per day = 84,000 / 21 = 4,000 ; OT rate = 84,000 / (21 x 8) x 1.5 = 750 / h
    absent  : Jun 2, Jun 3, Jun 10 (Jun 10 is also inside an approved unpaid leave -> counted once)
    half day: Jun 4
    unpaid leave with no attendance record: Jun 9
    LOP days = 3 + 0.5 + 1 = 4.5 -> deduction 18,000
    overtime = 60 + 60 minutes (Jun 5, Jun 6) -> 2 h x 750 = 1,500
    PF = 12 % x (84,000 x 35 % x 16.5 / 21) = 12 % x 23,100 = 2,772
    net = 84,000 + 1,500 - 2,772 - 18,000 = 64,728
"""

import io
import os
import sys
from datetime import date, datetime, time

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
from fastapi.testclient import TestClient

from app.database.connection import SessionLocal
from app.database.models import Attendance, Employee, Leave, Salary, User
from app.main import app
from app.services.employee_service import create_employee
from app.utils.security import create_access_token

client = TestClient(app)
db = SessionLocal()
passed_count = 0
failed_count = 0

CODES = ["T-PAY-MGR", "T-PAY-EMP", "T-PAY-NOSAL"]
DEPT = "QA Payroll"
PASSWORD = "PayTest#2026"


def chk(condition: bool, msg: str, fail_detail: str = ""):
    global passed_count, failed_count
    if condition:
        print(f"  [PASS] {msg}")
        passed_count += 1
    else:
        print(f"  [FAIL] {msg} -> {fail_detail}")
        failed_count += 1


def hdr(user: User):
    return {"Authorization": f"Bearer {create_access_token(user_id=user.id, role=user.role)}"}


def cleanup():
    db.rollback()
    emps = db.query(Employee).filter(Employee.employee_code.in_(CODES)).all()
    ids = [e.id for e in emps]
    if ids:
        db.query(Salary).filter(Salary.employee_id.in_(ids)).delete(synchronize_session=False)
        db.query(Attendance).filter(Attendance.employee_id.in_(ids)).delete(synchronize_session=False)
        db.query(Leave).filter(Leave.employee_id.in_(ids)).delete(synchronize_session=False)
        db.query(User).filter(User.employee_id.in_(ids)).delete(synchronize_session=False)
        db.query(Employee).filter(Employee.id.in_(ids)).update({Employee.manager_id: None}, synchronize_session=False)
        db.query(Employee).filter(Employee.id.in_(ids)).delete(synchronize_session=False)
    db.commit()


def att(emp_id, d, status, in_t=None, out_t=None, worked=0, ot=0):
    db.add(Attendance(employee_id=emp_id, attendance_date=d, status=status, in_time=in_t, out_time=out_t,
                      working_minutes=worked, late_minutes=0, overtime_minutes=ot))


def seeded_salary_fingerprint():
    db.rollback()
    rows = db.query(Salary).join(Employee, Employee.id == Salary.employee_id).filter(
        ~Employee.employee_code.in_(CODES)).order_by(Salary.id).all()
    return [(r.id, float(r.gross_salary), float(r.net_salary), r.paid_at) for r in rows]


try:
    cleanup()
    before_seeded = seeded_salary_fingerprint()

    hr = db.query(User).filter(User.email == "neha.hr@company.com").first()
    aman = db.query(User).filter(User.email == "aman@company.com").first()
    if not hr or not aman:
        raise RuntimeError("Seed users missing — run scripts/seed_db.py")

    mgr_emp = create_employee(db, employee_code="T-PAY-MGR", name="Pay Manager", department=DEPT,
                              designation="QA Lead", joining_date=date(2024, 1, 1),
                              email="t.pay.mgr@hrtest.dev", password=PASSWORD, role="manager")
    emp = create_employee(db, employee_code="T-PAY-EMP", name="Pay Employee", department=DEPT,
                          designation="QA Engineer", joining_date=date(2024, 1, 1), manager_id=mgr_emp.id,
                          email="t.pay.emp@hrtest.dev", password=PASSWORD, role="employee",
                          monthly_gross_salary=84000)
    nosal = create_employee(db, employee_code="T-PAY-NOSAL", name="Pay NoStructure", department=DEPT,
                            designation="QA Intern", joining_date=date(2024, 1, 1))
    mgr_user = db.query(User).filter(User.employee_id == mgr_emp.id).first()
    emp_user = db.query(User).filter(User.employee_id == emp.id).first()

    att(emp.id, date(2025, 6, 2), "absent")
    att(emp.id, date(2025, 6, 3), "absent")
    att(emp.id, date(2025, 6, 4), "half_day", time(9, 0), time(12, 0), 180)
    att(emp.id, date(2025, 6, 5), "present", time(9, 0), time(19, 0), 540, 60)
    att(emp.id, date(2025, 6, 6), "present", time(9, 0), time(19, 0), 540, 60)
    att(emp.id, date(2025, 6, 10), "absent")
    db.add(Leave(employee_id=emp.id, leave_type="unpaid", from_date=date(2025, 6, 9), to_date=date(2025, 6, 10),
                 status="approved", reason="T-PAY unpaid"))
    db.add(Leave(employee_id=emp.id, leave_type="unpaid", from_date=date(2025, 6, 16), to_date=date(2025, 6, 16),
                 status="pending", reason="T-PAY pending unpaid (must not count)"))
    db.commit()

    body = {"month": 6, "year": 2025, "employee_ids": [emp.id, nosal.id]}

    # ------------------------------------------------------------------
    print("\n[1] Auth & validation")
    chk(client.post("/salary/generate", json=body).status_code == 401, "Unauthenticated -> 401")
    chk(client.post("/salary/generate", json=body, headers=hdr(emp_user)).status_code == 403, "Employee -> 403")
    chk(client.post("/salary/generate", json=body, headers=hdr(mgr_user)).status_code == 403, "Manager -> 403")
    chk(client.post("/salary/generate", json={**body, "month": 13}, headers=hdr(hr)).status_code == 422, "Month 13 -> 422")
    future = date.today().replace(day=1)
    fy, fm = (future.year + 1, 1) if future.month == 12 else (future.year, future.month + 1)
    chk(client.post("/salary/generate", json={"month": fm, "year": fy, "employee_ids": [emp.id]},
                    headers=hdr(hr)).status_code == 400, "Future month -> 400")

    # ------------------------------------------------------------------
    print("\n[2] Happy path with hand-checked numbers")
    r = client.post("/salary/generate", json=body, headers=hdr(hr))
    data = r.json() if r.status_code == 200 else {}
    chk(r.status_code == 200 and data.get("created") == 1 and data.get("updated") == 0, "1 row created", r.text[:300])
    chk(data.get("working_days") == 21 and data.get("provisional") is False, "June 2025 = 21 working days, final", str(data.get("working_days")))
    item = (data.get("items") or [{}])[0]
    expected = {"gross_salary": 84000.0, "absent_days": 3, "half_days": 1, "unpaid_leave_days": 1, "lop_days": 4.5,
                "lop_deduction": 18000.0, "overtime_minutes": 120, "overtime_amount": 1500.0, "pf": 2772.0,
                "deductions": 18000.0, "net_salary": 64728.0}
    for key, val in expected.items():
        chk(item.get(key) == val, f"{key} = {val}", f"got {item.get(key)}")
    skipped = {s["employee_id"]: s["reason"] for s in data.get("skipped", [])}
    chk(nosal.id in skipped and "salary structure" in skipped[nosal.id], "Employee without salary structure is skipped with a reason",
        str(skipped))
    db.rollback()
    rows = db.query(Salary).filter(Salary.employee_id == emp.id, Salary.month == 6, Salary.year == 2025).all()
    chk(len(rows) == 1 and float(rows[0].net_salary) == 64728.0 and rows[0].paid_at is None, "Salary row persisted", str(rows))

    # ------------------------------------------------------------------
    print("\n[3] Idempotency")
    r = client.post("/salary/generate", json=body, headers=hdr(hr))
    chk(r.status_code == 200 and r.json()["created"] == 0 and r.json()["updated"] == 1, "Re-run updates, creates nothing", r.text[:200])
    db.rollback()
    chk(db.query(Salary).filter(Salary.employee_id == emp.id, Salary.month == 6, Salary.year == 2025).count() == 1,
        "Still exactly one row for the employee-month")

    att(emp.id, date(2025, 6, 11), "absent")
    db.commit()
    r = client.post("/salary/generate", json=body, headers=hdr(hr))
    item = r.json()["items"][0] if r.status_code == 200 else {}
    chk(item.get("lop_days") == 5.5 and item.get("net_salary") == round(84000 + 1500 - 0.12 * 29400 * 15.5 / 21 - 22000, 2),
        "New absence -> row recalculated (LOP 5.5)", str(item))

    db.rollback()
    row = db.query(Salary).filter(Salary.employee_id == emp.id, Salary.month == 6, Salary.year == 2025).first()
    row.paid_at = datetime(2025, 6, 30, 18, 0)
    net_before = float(row.net_salary)
    db.commit()
    att(emp.id, date(2025, 6, 12), "absent")
    db.commit()
    r = client.post("/salary/generate", json=body, headers=hdr(hr))
    sk = {s["employee_id"]: s["reason"] for s in r.json().get("skipped", [])} if r.status_code == 200 else {}
    chk(emp.id in sk and "locked" in sk[emp.id], "Paid row is locked and skipped", str(sk))
    db.rollback()
    chk(float(db.query(Salary).filter(Salary.id == row.id).first().net_salary) == net_before, "Paid row values unchanged")

    # ------------------------------------------------------------------
    print("\n[4] Ownership isolation & salary privacy")
    r = client.get("/salary/me", headers=hdr(emp_user))
    chk(r.status_code == 200 and any(s["month"] == 6 and s["year"] == 2025 for s in r.json()), "Employee sees own generated slip")
    r = client.get("/salary/me", headers=hdr(aman))
    chk(r.status_code == 200 and all(s["employee_id"] == aman.employee_id for s in r.json()), "Another employee never sees it")
    chk(client.get(f"/salary/{emp.id}", headers=hdr(mgr_user)).status_code == 403, "Manager cannot read team member salary -> 403")
    r = client.get(f"/employees/{emp.id}", headers=hdr(mgr_user))
    chk(r.status_code == 200 and r.json().get("monthly_gross_salary") is None, "Manager does not see salary structure", r.text[:200])
    r = client.get(f"/employees/{emp.id}", headers=hdr(emp_user))
    chk(r.status_code == 200 and r.json().get("monthly_gross_salary") == 84000.0, "Employee sees own salary structure")
    r = client.get(f"/employees/{emp.id}", headers=hdr(hr))
    chk(r.status_code == 200 and r.json().get("monthly_gross_salary") == 84000.0, "HR sees salary structure")
    r = client.put(f"/employees/{nosal.id}", json={"monthly_gross_salary": 30000}, headers=hdr(hr))
    chk(r.status_code == 200 and r.json().get("monthly_gross_salary") == 30000.0, "HR sets salary structure via PUT")
    r = client.post("/salary/generate", json={"month": 6, "year": 2025, "employee_ids": [nosal.id]}, headers=hdr(hr))
    chk(r.status_code == 200 and r.json()["created"] == 1 and r.json()["items"][0]["net_salary"] == round(30000 - 0.12 * 0.35 * 30000, 2),
        "Structure set -> payroll generated (no attendance = no LOP)", r.text[:200])

    # ------------------------------------------------------------------
    print("\n[5] KI-015: overtime amount pro-rated to the report period")
    r = client.get(f"/reports/overtime?date_from=2025-06-05&date_to=2025-06-05&department={DEPT}", headers=hdr(hr))
    wb = openpyxl.load_workbook(io.BytesIO(r.content)) if r.status_code == 200 else None
    row1 = next((x for x in wb["Summary"].iter_rows(min_row=5, values_only=True) if x and x[0] == "T-PAY-EMP"), None) if wb else None
    chk(row1 is not None and row1[3] == 1.0 and row1[4] == 750.0, "One day (60 min OT) -> 750.00", str(row1))
    r = client.get(f"/reports/overtime?month=6&year=2025&department={DEPT}", headers=hdr(hr))
    wb = openpyxl.load_workbook(io.BytesIO(r.content)) if r.status_code == 200 else None
    row2 = next((x for x in wb["Summary"].iter_rows(min_row=5, values_only=True) if x and x[0] == "T-PAY-EMP"), None) if wb else None
    chk(row2 is not None and row2[3] == 2.0 and row2[4] == 1500.0, "Whole month (120 min OT) -> 1,500.00 = payroll OT amount", str(row2))

    # ------------------------------------------------------------------
    print("\n[6] Seeded payroll untouched")
    chk(seeded_salary_fingerprint() == before_seeded, "No seeded salary row was created or modified")

except Exception as exc:
    import traceback
    traceback.print_exc()
    chk(False, "Unexpected exception", str(exc))
finally:
    cleanup()
    db.close()

print("\n" + "=" * 65)
print(f"  PAYROLL TEST RESULTS: {passed_count} PASSED, {failed_count} FAILED")
print("=" * 65)
sys.exit(1 if failed_count else 0)
