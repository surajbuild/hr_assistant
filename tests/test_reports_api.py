"""
tests/test_reports_api.py
-------------------------
Integration tests for Excel report generation & export:
  - GET /reports/attendance
  - GET /reports/overtime

Verifies:
1. Unauthorized request -> 401
2. Employee role -> 403
3. Manager role -> 403
4. HR role -> 200 with valid .xlsx stream & Content-Disposition
5. Admin role -> 200 with valid .xlsx stream & Content-Disposition
6. Date filtering (date_from and date_to)
7. Verification that returned bytes can be parsed by openpyxl
"""

import io
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
from fastapi.testclient import TestClient
from app.main import app
from app.database.connection import SessionLocal
from app.database.models import User
from app.utils.security import create_access_token

client = TestClient(app)
db = SessionLocal()


def get_token(role: str) -> str:
    user = db.query(User).filter(User.role == role).first()
    if not user:
        raise RuntimeError(f"User with role '{role}' not found in database.")
    return create_access_token(user_id=user.id, role=user.role)


emp_token = get_token("employee")
mgr_token = get_token("manager")
hr_token = get_token("hr")
admin_token = get_token("admin")

emp_hdr = {"Authorization": f"Bearer {emp_token}"}
mgr_hdr = {"Authorization": f"Bearer {mgr_token}"}
hr_hdr = {"Authorization": f"Bearer {hr_token}"}
admin_hdr = {"Authorization": f"Bearer {admin_token}"}

EXCEL_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ---------------------------------------------------------------------------
# Attendance Report Tests
# ---------------------------------------------------------------------------

def test_attendance_report_unauthorized():
    res = client.get("/reports/attendance")
    assert res.status_code == 401, f"Expected 401, got {res.status_code}"
    print("[PASS] test_attendance_report_unauthorized -> 401")


def test_attendance_report_employee_forbidden():
    res = client.get("/reports/attendance", headers=emp_hdr)
    assert res.status_code == 403, f"Expected 403, got {res.status_code}"
    print("[PASS] test_attendance_report_employee_forbidden -> 403")


def test_attendance_report_manager_forbidden():
    res = client.get("/reports/attendance", headers=mgr_hdr)
    assert res.status_code == 403, f"Expected 403, got {res.status_code}"
    print("[PASS] test_attendance_report_manager_forbidden -> 403")


def test_attendance_report_hr_allowed():
    res = client.get("/reports/attendance", headers=hr_hdr)
    assert res.status_code == 200, f"Expected 200, got {res.status_code}"
    assert EXCEL_CONTENT_TYPE in res.headers.get("content-type", "")
    assert "attachment; filename=" in res.headers.get("content-disposition", "")
    assert res.headers["content-disposition"].endswith('.xlsx"') or res.headers["content-disposition"].endswith('.xlsx')

    # Verify openpyxl can parse the returned bytes
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    assert "Attendance Records" in wb.sheetnames
    ws = wb["Attendance Records"]
    assert ws.max_row > 4  # Header rows + at least 1 record
    print(f"[PASS] test_attendance_report_hr_allowed -> 200 (rows: {ws.max_row})")


def test_attendance_report_admin_allowed():
    res = client.get("/reports/attendance", headers=admin_hdr)
    assert res.status_code == 200, f"Expected 200, got {res.status_code}"
    assert EXCEL_CONTENT_TYPE in res.headers.get("content-type", "")
    print("[PASS] test_attendance_report_admin_allowed -> 200")


def test_attendance_report_date_filter():
    res = client.get("/reports/attendance?date_from=2024-08-01&date_to=2024-08-31", headers=hr_hdr)
    assert res.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    ws = wb["Attendance Records"]
    # Check subtitle row contains the filtered date range
    sub_title = ws["A2"].value
    assert "2024-08-01" in sub_title and "2024-08-31" in sub_title
    print(f"[PASS] test_attendance_report_date_filter -> 200 (subtitle: {sub_title})")


# ---------------------------------------------------------------------------
# Overtime Report Tests
# ---------------------------------------------------------------------------

def test_overtime_report_unauthorized():
    res = client.get("/reports/overtime")
    assert res.status_code == 401, f"Expected 401, got {res.status_code}"
    print("[PASS] test_overtime_report_unauthorized -> 401")


def test_overtime_report_employee_forbidden():
    res = client.get("/reports/overtime", headers=emp_hdr)
    assert res.status_code == 403, f"Expected 403, got {res.status_code}"
    print("[PASS] test_overtime_report_employee_forbidden -> 403")


def test_overtime_report_manager_forbidden():
    res = client.get("/reports/overtime", headers=mgr_hdr)
    assert res.status_code == 403, f"Expected 403, got {res.status_code}"
    print("[PASS] test_overtime_report_manager_forbidden -> 403")


def test_overtime_report_hr_allowed():
    res = client.get("/reports/overtime", headers=hr_hdr)
    assert res.status_code == 200, f"Expected 200, got {res.status_code}"
    assert EXCEL_CONTENT_TYPE in res.headers.get("content-type", "")
    assert "attachment; filename=" in res.headers.get("content-disposition", "")

    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    assert "Overtime Records" in wb.sheetnames
    ws = wb["Overtime Records"]
    assert ws.max_row > 4  # Header rows + at least 1 record
    print(f"[PASS] test_overtime_report_hr_allowed -> 200 (rows: {ws.max_row})")


def test_overtime_report_admin_allowed():
    res = client.get("/reports/overtime", headers=admin_hdr)
    assert res.status_code == 200, f"Expected 200, got {res.status_code}"
    print("[PASS] test_overtime_report_admin_allowed -> 200")


def test_overtime_report_date_filter():
    res = client.get("/reports/overtime?date_from=2024-09-01&date_to=2024-09-30", headers=hr_hdr)
    assert res.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    ws = wb["Overtime Records"]
    sub_title = ws["A2"].value
    assert "2024-09-01" in sub_title and "2024-09-30" in sub_title
    print(f"[PASS] test_overtime_report_date_filter -> 200 (subtitle: {sub_title})")


if __name__ == "__main__":
    try:
        print("\n--- Running Attendance Report Tests ---")
        test_attendance_report_unauthorized()
        test_attendance_report_employee_forbidden()
        test_attendance_report_manager_forbidden()
        test_attendance_report_hr_allowed()
        test_attendance_report_admin_allowed()
        test_attendance_report_date_filter()

        print("\n--- Running Overtime Report Tests ---")
        test_overtime_report_unauthorized()
        test_overtime_report_employee_forbidden()
        test_overtime_report_manager_forbidden()
        test_overtime_report_hr_allowed()
        test_overtime_report_admin_allowed()
        test_overtime_report_date_filter()

        print("\n[SUCCESS] All 12 Reports API automated tests passed successfully!")
    finally:
        db.close()
